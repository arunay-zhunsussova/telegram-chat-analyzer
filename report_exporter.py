from __future__ import annotations

from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEET_ORDER = ["overview", "manager_summary", "personal_recommendations", "mood_flow", "topics", "critical_cases", "messages"]


def value_to_cell(value: Any) -> Any:
    if value is None:
        return ""
    return value


def rows_to_sheet(wb: Workbook, sheet_name: str, rows: List[Dict[str, Any]]) -> None:
    ws = wb.create_sheet(sheet_name)
    if rows:
        headers = list(rows[0].keys())
    else:
        headers = []
    ws.append(headers)
    for row in rows:
        ws.append([value_to_cell(row.get(header, "")) for header in headers])

    ws.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for idx, header in enumerate(headers, start=1):
        values = [str(row.get(header, "")) for row in rows]
        max_len = max([len(str(header))] + [len(v) for v in values])
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 14), 60)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def reports_to_excel_bytes(reports: Dict[str, List[Dict[str, Any]]]) -> bytes:
    wb = Workbook()
    default = wb.active
    wb.remove(default)
    for sheet_name in SHEET_ORDER:
        rows_to_sheet(wb, sheet_name, reports.get(sheet_name, []))
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def save_reports_to_excel(reports: Dict[str, List[Dict[str, Any]]], output_path: str | Path) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_bytes(reports_to_excel_bytes(reports))
    return output_path
