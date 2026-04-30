from __future__ import annotations

import csv
import json
import math
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook

SUPPORTED_SUFFIXES = {".json", ".csv", ".xlsx", ".txt", ".md"}
REPORT_SHEETS = ["overview", "manager_summary", "personal_recommendations", "mood_flow", "topics", "critical_cases", "messages"]

POSITIVE_WORDS = {
    "thank", "thanks", "good", "great", "perfect", "helpful", "approved", "sounds good",
    "happy", "excellent", "sure", "okay", "ok", "reserved", "confirmed", "understood",
}
NEGATIVE_WORDS = {
    "rude", "bad", "problem", "delay", "delayed", "cancel", "broken", "not work", "does not turn on",
    "disappointed", "complaint", "bad review", "not clear", "confused", "too blue", "harsh", "angry",
    "refund", "exchange", "quality", "not my problem", "stupid", "do not care", "whatever",
}
AGGRESSIVE_WORDS = {"stupid", "do not care", "don't care", "whatever", "not my problem", "stop acting", "do what you want"}
HELPFUL_WORDS = {"happy to help", "recommend", "suggest", "sorry", "apologize", "please", "thank", "i can help", "confirmed", "understood"}
DISMISSIVE_SHORT = {"sure", "maybe", "okay", "ok", "check it", "looks fine", "i do not know", "i don't know"}

TOPIC_KEYWORDS: Dict[str, Sequence[str]] = {
    "lighting color": ["warm", "cold", "2700k", "3000k", "4000k", "6500k", "yellow", "blue", "color"],
    "bulb/lamp choice": ["lamp", "bulb", "led", "watt", "watts", "e27", "e14", "socket", "fit"],
    "smart bulb setup": ["smart", "app", "wi-fi", "wifi", "bluetooth", "setup", "pair", "device"],
    "delivery/order": ["delivery", "courier", "order", "delayed", "delay", "arrives", "reserve", "pick it up"],
    "refund/warranty": ["refund", "replacement", "warranty", "broken", "does not turn on", "exchange", "quality"],
    "price/payment": ["price", "payment", "invoice", "discount", "bulk", "tenge", "total"],
    "service quality": ["rude", "bad review", "not clear", "disappointed", "not my problem", "stupid", "do not care"],
    "energy saving": ["energy", "saving", "electricity", "last longer", "fluorescent"],
}


def safe_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and math.isnan(value):
        return ""
    if isinstance(value, list):
        parts: List[str] = []
        for item in value:
            if isinstance(item, str):
                parts.append(item)
            elif isinstance(item, dict):
                parts.append(str(item.get("text", "")))
            else:
                parts.append(str(item))
        return "".join(parts).strip()
    if isinstance(value, dict):
        return str(value.get("text", value)).strip()
    return str(value).strip()


def parse_dt(value: Any) -> Optional[datetime]:
    text = safe_str(value)
    if not text:
        return None
    if isinstance(value, datetime):
        return value
    for fmt in [
        "%Y-%m-%dT%H:%M:%S",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%d.%m.%Y %H:%M",
        "%d/%m/%Y %H:%M",
        "%d-%m-%Y %H:%M",
    ]:
        try:
            return datetime.strptime(text[:19], fmt)
        except ValueError:
            pass
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
    except Exception:
        return None


def dt_to_str(value: Any) -> str:
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    return safe_str(value)


def contains_any(text: str, words: Iterable[str]) -> bool:
    low = text.lower()
    return any(word in low for word in words)


def detect_role(sender: str, role: str = "", manager_hints: Optional[Sequence[str]] = None) -> str:
    role_low = role.strip().lower()
    if role_low in {"manager", "client"}:
        return role_low
    sender_low = sender.lower()
    hints = [h.strip().lower() for h in (manager_hints or []) if h.strip()]
    if hints and any(h in sender_low for h in hints):
        return "manager"
    if any(w in sender_low for w in ["manager", "support", "sales", "operator", "admin", "consultant"]):
        return "manager"
    return "client"


def auto_column(columns: Sequence[str], candidates: Sequence[str]) -> Optional[str]:
    low_map = {str(c).lower().strip(): str(c) for c in columns}
    for cand in candidates:
        if cand.lower() in low_map:
            return low_map[cand.lower()]
    for c in columns:
        low = str(c).lower().strip()
        if any(cand.lower() in low for cand in candidates):
            return str(c)
    return None


@dataclass
class InputOptions:
    input_format: str = "auto"
    sender_col: str = ""
    text_col: str = ""
    timestamp_col: str = ""
    role_col: str = ""
    message_id_col: str = ""
    chat_title_col: str = ""
    manager_hints: Tuple[str, ...] = ()
    window_size: int = 6


class InputFormatError(ValueError):
    pass


def normalize_file(path: str | Path, options: Optional[InputOptions] = None) -> List[Dict[str, Any]]:
    options = options or InputOptions()
    path = Path(path)
    fmt = options.input_format
    if fmt == "auto":
        suffix = path.suffix.lower()
        if suffix == ".json":
            fmt = "telegram_json"
        elif suffix == ".csv":
            fmt = "csv_table"
        elif suffix in {".xlsx"}:
            fmt = "excel_table"
        elif suffix in {".txt", ".md"}:
            fmt = "plain_text"
        else:
            raise InputFormatError(f"Unsupported file type: {path.suffix}")

    if fmt == "telegram_json":
        return normalize_telegram_json(path, options)
    if fmt == "csv_table":
        return normalize_table(read_csv_rows(path), path.name, options)
    if fmt == "excel_table":
        return normalize_table(read_excel_rows(path), path.name, options)
    if fmt == "plain_text":
        return normalize_plain_text(path.read_text(encoding="utf-8", errors="ignore"), path.name, options)
    raise InputFormatError(f"Unknown input format: {options.input_format}")


def normalize_telegram_json(path: str | Path, options: InputOptions) -> List[Dict[str, Any]]:
    path = Path(path)
    with path.open("r", encoding="utf-8") as f:
        data = json.load(f)
    chat_title = safe_str(data.get("name")) or path.stem
    rows: List[Dict[str, Any]] = []
    for index, msg in enumerate(data.get("messages", []), start=1):
        if not isinstance(msg, dict):
            continue
        if msg.get("type") not in {None, "message", "service"}:
            continue
        text = safe_str(msg.get("text"))
        if not text:
            continue
        sender = safe_str(msg.get("from")) or safe_str(msg.get("actor")) or "Unknown"
        rows.append(
            {
                "source_file": path.name,
                "chat_title": chat_title,
                "message_id": msg.get("id", index),
                "timestamp": parse_dt(msg.get("date")),
                "sender": sender,
                "role": detect_role(sender, manager_hints=options.manager_hints),
                "text": text,
            }
        )
    return rows


def read_csv_rows(path: str | Path) -> List[Dict[str, Any]]:
    path = Path(path)
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def read_excel_rows(path: str | Path) -> List[Dict[str, Any]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [safe_str(h) or f"column_{i+1}" for i, h in enumerate(rows[0])]
    output: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        output.append({headers[i]: raw[i] if i < len(raw) else "" for i in range(len(headers))})
    return output


def normalize_table(rows: List[Dict[str, Any]], source_file: str, options: InputOptions) -> List[Dict[str, Any]]:
    if not rows:
        return []
    columns = list(rows[0].keys())
    sender_col = options.sender_col or auto_column(columns, ["sender", "from", "name", "author", "user"])
    text_col = options.text_col or auto_column(columns, ["text", "message", "content", "body"])
    timestamp_col = options.timestamp_col or auto_column(columns, ["timestamp", "date", "datetime", "time", "created_at"])
    role_col = options.role_col or auto_column(columns, ["role"])
    message_id_col = options.message_id_col or auto_column(columns, ["message_id", "id", "msg_id"])
    chat_title_col = options.chat_title_col or auto_column(columns, ["chat_title", "chat", "conversation", "dialog"])
    if text_col is None:
        raise InputFormatError("Could not find a text/message column. Set it in Field mapping.")
    if sender_col is None:
        raise InputFormatError("Could not find a sender/from column. Set it in Field mapping.")

    output: List[Dict[str, Any]] = []
    for i, row in enumerate(rows, start=1):
        text = safe_str(row.get(text_col))
        if not text:
            continue
        sender = safe_str(row.get(sender_col)) or "Unknown"
        role_raw = safe_str(row.get(role_col)) if role_col else ""
        output.append(
            {
                "source_file": source_file,
                "chat_title": safe_str(row.get(chat_title_col)) if chat_title_col else Path(source_file).stem,
                "message_id": row.get(message_id_col, i) if message_id_col else i,
                "timestamp": parse_dt(row.get(timestamp_col)) if timestamp_col else None,
                "sender": sender,
                "role": detect_role(sender, role_raw, options.manager_hints),
                "text": text,
            }
        )
    return output


TEXT_LINE_PATTERNS = [
    re.compile(r"^(?P<timestamp>\d{4}-\d{2}-\d{2}[ T]\d{1,2}:\d{2}(?::\d{2})?)\s*[-–]\s*(?P<sender>[^:]+):\s*(?P<text>.+)$"),
    re.compile(r"^(?P<timestamp>\d{1,2}[./-]\d{1,2}[./-]\d{2,4}\s+\d{1,2}:\d{2})\s*[-–]\s*(?P<sender>[^:]+):\s*(?P<text>.+)$"),
    re.compile(r"^(?P<sender>[^:]{1,80}):\s*(?P<text>.+)$"),
]


def normalize_plain_text(text: str, source_file: str, options: InputOptions) -> List[Dict[str, Any]]:
    rows: List[Dict[str, Any]] = []
    current_sender = "Unknown"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        match = None
        for pattern in TEXT_LINE_PATTERNS:
            match = pattern.match(line)
            if match:
                break
        if match:
            gd = match.groupdict()
            sender = safe_str(gd.get("sender")) or current_sender
            current_sender = sender
            msg_text = safe_str(gd.get("text"))
            timestamp = parse_dt(gd.get("timestamp")) if gd.get("timestamp") else None
        else:
            sender = current_sender
            msg_text = line
            timestamp = None
        if not msg_text:
            continue
        rows.append(
            {
                "source_file": source_file,
                "chat_title": Path(source_file).stem,
                "message_id": len(rows) + 1,
                "timestamp": timestamp,
                "sender": sender,
                "role": detect_role(sender, manager_hints=options.manager_hints),
                "text": msg_text,
            }
        )
    return rows


def mood_score(text: str) -> int:
    low = text.lower()
    score = sum(1 for word in POSITIVE_WORDS if word in low)
    score -= sum(1 for word in NEGATIVE_WORDS if word in low)
    return score


def label_client_mood(client_text: str) -> str:
    low = client_text.lower()
    score = mood_score(client_text)
    if any(w in low for w in ["bad review", "rude", "angry", "stupid"]):
        return "angry"
    if any(w in low for w in ["confused", "not understand", "not sure", "which", "what about"]):
        return "confused"
    if score <= -2:
        return "frustrated"
    if score >= 2:
        return "positive"
    return "neutral"


def label_manager_tone(manager_text: str) -> str:
    low = manager_text.lower().strip()
    if contains_any(low, AGGRESSIVE_WORDS):
        return "aggressive"
    if low in DISMISSIVE_SHORT or (len(low.split()) <= 3 and low in {"sure", "maybe", "okay", "ok"}):
        return "dismissive"
    if contains_any(low, HELPFUL_WORDS):
        return "helpful"
    return "neutral"


def extract_topics(text: str) -> List[str]:
    low = text.lower()
    topics = [topic for topic, words in TOPIC_KEYWORDS.items() if any(word in low for word in words)]
    return topics or ["general support"]


def risk_signals(client_text: str, manager_text: str) -> List[str]:
    combined = f"{client_text}\n{manager_text}".lower()
    risks: List[str] = []
    if any(w in combined for w in ["bad review", "cancel", "disappointed", "rude", "churn"]):
        risks.append("churn risk")
    if any(w in manager_text.lower() for w in AGGRESSIVE_WORDS):
        risks.append("aggressive manager tone")
    if any(w in combined for w in ["refund", "replacement", "exchange", "broken", "warranty"]):
        risks.append("refund/warranty issue")
    if any(w in combined for w in ["delay", "delayed", "courier", "promised yesterday"]):
        risks.append("delivery delay")
    if any(w in manager_text.lower() for w in ["maybe", "i do not know", "i don't know", "sure.", "check it"]):
        risks.append("unclear answer")
    return risks or ["none"]


def response_minutes(messages: List[Dict[str, Any]]) -> List[float]:
    sorted_messages = sorted(messages, key=lambda m: (m.get("timestamp") is None, m.get("timestamp") or datetime.max, int_safe(m.get("message_id"))))
    times: List[float] = []
    last_client_time: Optional[datetime] = None
    for msg in sorted_messages:
        ts = msg.get("timestamp")
        if not isinstance(ts, datetime):
            continue
        if msg.get("role") == "client":
            last_client_time = ts
        elif msg.get("role") == "manager" and last_client_time is not None:
            delta = (ts - last_client_time).total_seconds() / 60
            if delta >= 0:
                times.append(round(delta, 2))
            last_client_time = None
    return times


def int_safe(value: Any) -> int:
    try:
        return int(value)
    except Exception:
        return 0


def trend_from_scores(scores: List[int]) -> str:
    if len(scores) < 2:
        return "stable"
    midpoint = max(1, len(scores) // 2)
    first = sum(scores[:midpoint])
    second = sum(scores[midpoint:])
    if second - first >= 2:
        return "improving"
    if first - second >= 2:
        return "worsening"
    return "stable"


def group_by_source(messages: List[Dict[str, Any]]) -> Dict[str, List[Dict[str, Any]]]:
    groups: Dict[str, List[Dict[str, Any]]] = {}
    for msg in messages:
        groups.setdefault(safe_str(msg.get("source_file")) or "chat", []).append(msg)
    return groups


def analyze_chat(messages: List[Dict[str, Any]], window_size: int = 6) -> Dict[str, List[Dict[str, Any]]]:
    reports: Dict[str, List[Dict[str, Any]]] = {sheet: [] for sheet in REPORT_SHEETS}
    if not messages:
        return reports

    for source_file, chat_messages in group_by_source(messages).items():
        chat_messages = sorted(chat_messages, key=lambda m: (m.get("timestamp") is None, m.get("timestamp") or datetime.max, int_safe(m.get("message_id"))))
        chat_title = safe_str(chat_messages[0].get("chat_title")) or source_file
        manager_names = sorted({safe_str(m.get("sender")) for m in chat_messages if m.get("role") == "manager" and safe_str(m.get("sender"))})
        manager_name = ", ".join(manager_names) if manager_names else "Not detected"
        total = len(chat_messages)
        client_count = sum(1 for m in chat_messages if m.get("role") == "client")
        manager_count = sum(1 for m in chat_messages if m.get("role") == "manager")
        resp_times = response_minutes(chat_messages)
        avg_response = round(sum(resp_times) / len(resp_times), 2) if resp_times else ""

        all_text = "\n".join(safe_str(m.get("text")) for m in chat_messages)
        topic_counts: Dict[str, int] = {}
        for msg in chat_messages:
            for topic in extract_topics(safe_str(msg.get("text"))):
                topic_counts[topic] = topic_counts.get(topic, 0) + 1
        for topic, count in sorted(topic_counts.items(), key=lambda item: (-item[1], item[0])):
            reports["topics"].append(
                {
                    "source_file": source_file,
                    "chat_title": chat_title,
                    "manager_name": manager_name,
                    "topic": topic,
                    "mentions": count,
                }
            )

        mood_scores: List[int] = []
        risk_count = 0
        helpful_count = 0
        aggressive_count = 0
        dismissive_count = 0
        windows = [chat_messages[i : i + max(1, window_size)] for i in range(0, len(chat_messages), max(1, window_size))]
        for win_idx, window in enumerate(windows, start=1):
            client_text = "\n".join(safe_str(m.get("text")) for m in window if m.get("role") == "client")
            manager_text = "\n".join(safe_str(m.get("text")) for m in window if m.get("role") == "manager")
            client_mood = label_client_mood(client_text)
            manager_tone = label_manager_tone(manager_text)
            helpful_count += 1 if manager_tone == "helpful" else 0
            aggressive_count += 1 if manager_tone == "aggressive" else 0
            dismissive_count += 1 if manager_tone == "dismissive" else 0
            risks = risk_signals(client_text, manager_text)
            if any(r != "none" for r in risks):
                risk_count += 1
            score = mood_score(client_text + "\n" + manager_text)
            mood_scores.append(score)
            topics = sorted({topic for msg in window for topic in extract_topics(safe_str(msg.get("text")))})
            note = make_short_note(client_mood, manager_tone, risks, topics)
            reports["mood_flow"].append(
                {
                    "source_file": source_file,
                    "chat_title": chat_title,
                    "manager_name": manager_name,
                    "window_id": win_idx,
                    "client_mood": client_mood,
                    "client_mood_score": score,
                    "manager_tone": manager_tone,
                    "conversation_trend": trend_from_scores(mood_scores),
                    "risk_signals": ", ".join(risks),
                    "topics": ", ".join(topics),
                    "short_note": note,
                }
            )
            if any(r != "none" for r in risks):
                reports["critical_cases"].append(
                    {
                        "source_file": source_file,
                        "chat_title": chat_title,
                        "manager_name": manager_name,
                        "window_id": win_idx,
                        "manager_tone": manager_tone,
                        "risk_signals": ", ".join(risks),
                        "note": note,
                    }
                )

        avg_response_number = float(avg_response) if avg_response != "" else 0.0
        avg_response_penalty = min(25, int(avg_response_number * 2))
        risk_penalty = min(45, risk_count * 15 + aggressive_count * 20 + dismissive_count * 8)
        helpful_bonus = min(15, helpful_count * 5)
        efficiency_score = max(0, min(100, 80 - avg_response_penalty - risk_penalty + helpful_bonus))
        final_verdict = verdict(efficiency_score, risk_count, aggressive_count)
        pain_points = repeated_pain_points(all_text)
        missed = missed_opportunities(all_text, helpful_count, risk_count)
        strengths = strengths_from_analysis(helpful_count, avg_response_number if avg_response != "" else None, risk_count, all_text)
        actions = improvement_actions(risk_count, aggressive_count, dismissive_count, avg_response_number if avg_response != "" else None)
        top_topics = [topic for topic, _ in sorted(topic_counts.items(), key=lambda item: (-item[1], item[0]))[:3]]
        recommendation = build_personal_recommendation(
            manager_name=manager_name,
            score=efficiency_score,
            final_verdict=final_verdict,
            risk_count=risk_count,
            aggressive_count=aggressive_count,
            dismissive_count=dismissive_count,
            helpful_count=helpful_count,
            avg_response=avg_response_number if avg_response != "" else None,
            top_topics=top_topics,
            pain_points=pain_points,
            missed_opportunities_list=missed,
            strengths=strengths,
            actions=actions,
        )

        reports["overview"].append(
            {
                "source_file": source_file,
                "chat_title": chat_title,
                "manager_name": manager_name,
                "windows": len(windows),
                "unique_topics": len(topic_counts),
                "total_messages": total,
                "client_messages": client_count,
                "manager_messages": manager_count,
                "average_response_minutes": avg_response,
                "efficiency_score": efficiency_score,
                "final_verdict": final_verdict,
            }
        )
        reports["manager_summary"].append(
            {
                "source_file": source_file,
                "manager_name": manager_name,
                "chat_title": chat_title,
                "total_messages": total,
                "client_messages": client_count,
                "manager_messages": manager_count,
                "average_response_minutes": avg_response,
                "communication_quality": communication_quality(efficiency_score, aggressive_count, dismissive_count),
                "empathy_and_politeness": empathy_label(aggressive_count, dismissive_count, helpful_count),
                "sales_or_support_effectiveness": sales_effectiveness(all_text, risk_count, helpful_count),
                "repeated_pain_points": ", ".join(pain_points),
                "missed_opportunities": ", ".join(missed),
                "strengths": ", ".join(strengths),
                "improvement_actions": ", ".join(actions),
                "personal_recommendation": recommendation["personal_recommendation"],
                "efficiency_score": efficiency_score,
                "final_verdict": final_verdict,
            }
        )
        reports["personal_recommendations"].append(
            {
                "source_file": source_file,
                "manager_name": manager_name,
                "chat_title": chat_title,
                "priority_level": recommendation["priority_level"],
                "main_focus": recommendation["main_focus"],
                "strength_to_keep": recommendation["strength_to_keep"],
                "personal_recommendation": recommendation["personal_recommendation"],
                "coaching_plan": recommendation["coaching_plan"],
                "example_phrase": recommendation["example_phrase"],
                "efficiency_score": efficiency_score,
                "final_verdict": final_verdict,
            }
        )

    for msg in messages:
        clean = dict(msg)
        clean["timestamp"] = dt_to_str(clean.get("timestamp"))
        reports["messages"].append(clean)
    return reports


def make_short_note(client_mood: str, manager_tone: str, risks: Sequence[str], topics: Sequence[str]) -> str:
    topic_text = ", ".join(list(topics)[:3]) if topics else "general support"
    if any(r != "none" for r in risks):
        return f"Risk detected: {', '.join(risks)}. Topic: {topic_text}."
    return f"Client mood is {client_mood}; manager tone is {manager_tone}. Topic: {topic_text}."


def repeated_pain_points(text: str) -> List[str]:
    low = text.lower()
    checks = {
        "delivery delay": ["delay", "delayed", "courier", "promised yesterday"],
        "unclear consultation": ["not understand", "not sure", "not clear", "confused"],
        "refund or warranty issue": ["refund", "warranty", "replacement", "broken"],
        "wrong lighting recommendation": ["too blue", "harsh", "cold", "warm light"],
        "rude support tone": ["rude", "stupid", "not my problem", "do not care"],
    }
    points = [label for label, words in checks.items() if any(w in low for w in words)]
    return points or ["No repeated pain point detected"]


def missed_opportunities(text: str, helpful_count: int, risk_count: int) -> List[str]:
    low = text.lower()
    missed: List[str] = []
    if "can i send a photo" in low and "looks fine" in low:
        missed.append("Ask clarifying photo/socket questions instead of vague answer")
    if any(w in low for w in ["refund", "warranty", "broken"]) and "how long" in low:
        missed.append("Explain refund/warranty steps and timing clearly")
    if risk_count and helpful_count == 0:
        missed.append("Use apology, ownership, and clear next step in risky conversations")
    if any(w in low for w in ["bulk", "invoice", "new apartment", "bundle"]) and "discount" not in low:
        missed.append("Offer bundle/discount or next-step invoice")
    return missed or ["No major missed opportunity detected"]


def strengths_from_analysis(helpful_count: int, avg_response: Optional[float], risk_count: int, text: str) -> List[str]:
    low = text.lower()
    strengths: List[str] = []
    if helpful_count > 0:
        strengths.append("Helpful recommendations")
    if avg_response is not None and avg_response <= 5:
        strengths.append("Fast response time")
    if "recommend" in low or "suggest" in low:
        strengths.append("Product-fit explanation")
    if "warranty" in low or "discount" in low or "invoice" in low:
        strengths.append("Commercial details provided")
    if risk_count == 0:
        strengths.append("Low risk conversation")
    return strengths or ["Basic response coverage"]


def improvement_actions(risk_count: int, aggressive_count: int, dismissive_count: int, avg_response: Optional[float]) -> List[str]:
    actions: List[str] = []
    if aggressive_count:
        actions.append("Remove rude phrases and use polite conflict recovery")
    if dismissive_count:
        actions.append("Replace one-word replies with clear explanations")
    if risk_count:
        actions.append("Give concrete next step for complaint/refund/delivery cases")
    if avg_response is not None and avg_response > 10:
        actions.append("Reduce response time for client questions")
    return actions or ["Keep asking clarifying questions and giving specific recommendations"]


def build_personal_recommendation(
    manager_name: str,
    score: int,
    final_verdict: str,
    risk_count: int,
    aggressive_count: int,
    dismissive_count: int,
    helpful_count: int,
    avg_response: Optional[float],
    top_topics: Sequence[str],
    pain_points: Sequence[str],
    missed_opportunities_list: Sequence[str],
    strengths: Sequence[str],
    actions: Sequence[str],
) -> Dict[str, str]:
    """Create a human-readable recommendation from the rule-based analysis results."""
    name = manager_name if manager_name and manager_name != "Not detected" else "this manager"
    topics = ", ".join(top_topics) if top_topics else "general support"
    main_pain = next((p for p in pain_points if not p.lower().startswith("no ")), "No repeated pain point detected")
    main_action = actions[0] if actions else "Keep giving specific recommendations"
    strength = strengths[0] if strengths else "Basic response coverage"

    if aggressive_count:
        priority = "High priority"
        focus = "Conflict recovery and polite language"
        recommendation = (
            f"{name} needs immediate coaching on polite conflict handling. The analysis found aggressive or risky language, "
            f"so the manager should avoid defensive phrases, apologize clearly, and give one concrete next step. "
            f"Main topic area: {topics}."
        )
        coaching_plan = (
            "1) Remove rude or dismissive phrases. 2) Start complaint replies with apology and ownership. "
            "3) End every risky conversation with a clear action, deadline, or responsible person."
        )
        example_phrase = "Sorry for the inconvenience. I understand the issue, and I will check it now and give you the next step."
    elif risk_count:
        priority = "Medium-high priority"
        focus = "Risk handling and clear next steps"
        recommendation = (
            f"{name} should improve risk recovery. The chat contains warning signs such as {main_pain.lower()}, "
            f"so the manager should respond with more ownership, empathy, and exact next steps. "
            f"Main topic area: {topics}."
        )
        coaching_plan = (
            f"Focus on this action: {main_action}. The manager should explain what will happen next, "
            "when the client will get an update, and what options are available."
        )
        example_phrase = "I understand your concern. Here is what we can do now: first, I will check the details, then I will offer the best solution."
    elif dismissive_count:
        priority = "Medium priority"
        focus = "Detailed answers and empathy"
        recommendation = (
            f"{name} should make answers more complete. The manager's replies look too short or unclear in some parts, "
            f"so clients may feel that their problem is not fully handled. Main topic area: {topics}."
        )
        coaching_plan = (
            "Use a 3-part answer: acknowledge the question, give the recommendation, and explain why. "
            "Avoid one-word answers when the client is asking for help or advice."
        )
        example_phrase = "Yes, this option can work for you because it matches your request. I also recommend checking this detail before buying."
    elif avg_response is not None and avg_response > 10:
        priority = "Medium priority"
        focus = "Response speed"
        recommendation = (
            f"{name} communicates safely, but response time should be improved. The average response time is {avg_response} minutes, "
            "which can reduce client satisfaction even when the answer is correct."
        )
        coaching_plan = "Try to send a short first reply quickly, then provide the full answer after checking the details."
        example_phrase = "I received your question. I am checking the details and will send you the exact answer shortly."
    elif score >= 80 and helpful_count > 0:
        priority = "Low priority"
        focus = "Maintain strong consultation style"
        recommendation = (
            f"{name} shows strong consultation quality. The manager should keep the current style: helpful tone, clear recommendations, "
            f"and low-risk communication. Main strength: {strength.lower()}."
        )
        coaching_plan = "Maintain the same approach and add small upsell or follow-up questions when relevant."
        example_phrase = "Based on your situation, I recommend this option. It fits your needs because of these features."
    else:
        priority = "Normal priority"
        focus = "More proactive consultation"
        recommendation = (
            f"{name} gives basic support, but the conversation can become more useful. The manager should ask more clarifying questions, "
            f"give specific recommendations, and close the conversation with a clear next step. Main topic area: {topics}."
        )
        coaching_plan = f"Focus on this action: {main_action}. Also keep this strength: {strength.lower()}."
        example_phrase = "To recommend the best option, may I ask one more question about your goal or situation?"

    return {
        "priority_level": priority,
        "main_focus": focus,
        "strength_to_keep": strength,
        "personal_recommendation": recommendation,
        "coaching_plan": coaching_plan,
        "example_phrase": example_phrase,
        "final_verdict": final_verdict,
    }


def verdict(score: int, risk_count: int, aggressive_count: int) -> str:
    if aggressive_count:
        return "Critical: aggressive manager behavior"
    if score >= 80 and risk_count == 0:
        return "Excellent consultation"
    if score >= 60:
        return "Good, but can improve"
    if score >= 40:
        return "Weak consultation"
    return "Critical service risk"


def communication_quality(score: int, aggressive_count: int, dismissive_count: int) -> str:
    if aggressive_count:
        return "Poor: aggressive language detected"
    if dismissive_count:
        return "Weak: answers are too short or unclear"
    if score >= 75:
        return "Strong: clear and useful communication"
    if score >= 50:
        return "Medium: acceptable but needs more detail"
    return "Weak: client needs are not handled well"


def empathy_label(aggressive_count: int, dismissive_count: int, helpful_count: int) -> str:
    if aggressive_count:
        return "Very low"
    if dismissive_count:
        return "Low"
    if helpful_count:
        return "Good"
    return "Neutral"


def sales_effectiveness(text: str, risk_count: int, helpful_count: int) -> str:
    low = text.lower()
    if risk_count:
        return "Low in risky parts; needs recovery strategy"
    if any(w in low for w in ["bundle", "discount", "invoice", "approved", "payment link", "reserved"]):
        return "High: moved conversation toward order or clear next step"
    if helpful_count:
        return "Good: recommendations are present"
    return "Medium: basic answers only"


def analyze_files(paths: Sequence[str | Path], options: Optional[InputOptions] = None) -> Dict[str, List[Dict[str, Any]]]:
    options = options or InputOptions()
    all_messages: List[Dict[str, Any]] = []
    for path in paths:
        all_messages.extend(normalize_file(path, options))
    return analyze_chat(all_messages, options.window_size)
